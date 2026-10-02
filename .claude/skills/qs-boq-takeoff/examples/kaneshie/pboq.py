from qty import Q,DOORS,WINS,SHEETS
from openpyxl import load_workbook
from openpyxl.styles import Font,Alignment,Border,Side,PatternFill
import copy
wb=load_workbook('fmt.xlsx')
FN='Calibri'; SZ=18
med=Side(style='medium'); dbl=Side(style='double'); thin=Side(style='thin')
NUM='_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
QNUM='_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
# ================= RATES =================
DW={n:(w,h) for n,w,h,q in DOORS+WINS}
RATES=[('S','SUBSTRUCTURE / EARTHWORKS',None,None,None),
('R_CLR','Clear site of vegetation, grub up roots and dispose','m2',15,'Machine & labour'),
('R_TOP','Excavate oversite avg 150mm deep, deposit on site','m2',28,''),
('R_EXC','Excavate to reduce levels max 4.00m deep (machine)','m3',85,''),
('R_FILL1','Imported fill over raft incl. membrane/geogrid, compacted','m3',220,''),
('R_FILL2','Crushed rock fill in 300mm compacted layers with geogrid','m3',480,''),
('R_BACK','Backfill with selected excavated material, compacted','m3',60,''),
('R_DISP','Load and cart away surplus material','m3',75,''),
('R_HARD','Hardcore filling 150mm, compacted','m3',260,''),
('R_BLIND','Plain concrete C15 blinding','m3',1650,''),
('R_BED','Concrete bed 150mm','m3',2150,''),
('C','IN-SITU CONCRETE C40 (supply, place, vibrate, cure)',None,None,None),
('R_CRAFT','C40 concrete in raft','m3',2550,''),
('R_CWALL','C40 concrete in walls','m3',2750,''),
('R_CCOL','C40 concrete in columns','m3',2800,''),
('R_CSLAB','C40 concrete in suspended slabs','m3',2650,''),
('R_CBEAM','C40 concrete in beams','m3',2700,''),
('R_CSTAIR','C40 concrete in stairs','m3',2850,''),
('R_CLINT','Concrete in lintels','m3',2600,''),
('R_CCOP','Concrete in copings','m3',2600,''),
('T','REINFORCEMENT (supply, cut, bend, fix incl. tying wire)',None,None,None),
('R_Y10','High yield bar 10mm','t',20500,'GHS per tonne'),
('R_Y12','High yield bar 12mm','t',20000,''),
('R_Y16','High yield bar 16mm','t',19500,''),
('R_Y20','High yield bar 20mm','t',19000,''),
('R_Y25','High yield bar 25mm','t',19000,''),
('F','FORMWORK',None,None,None),
('R_FWALL','Formwork to walls','m2',185,''),
('R_FCOL','Formwork to columns','m2',190,''),
('R_FSLAB','Formwork to slab soffits incl. props','m2',170,''),
('R_FBEAM','Formwork to beam sides & soffits','m2',200,''),
('R_FEDGE','Formwork to edges not exceeding 250mm high','m',65,''),
('R_FRAFT','Formwork to raft edges','m2',150,''),
('R_FSTAIR','Formwork to stairs','m2',230,''),
('R_FLINT','Formwork to lintels / copings','m2',180,''),
('M','MASONRY & FINISHES',None,None,None),
('R_B150','150mm sandcrete blockwork','m2',215,''),
('R_B200','200mm sandcrete blockwork','m2',265,''),
('R_REND','12.5mm cement & sand rendering','m2',78,''),
('R_RENDX','External rendering','m2',88,''),
('R_SCR25','25mm screed','m2',68,''),
('R_BED12','12mm screeded bed to receive floor tiles','m2',48,''),
('R_BACK12','12mm backing to receive wall tiles','m2',55,''),
('R_T40','400x400 non-slip porcelain tiles','m2',230,'Supply & fix'),
('R_T50','500x500 matte R11 porcelain tiles to stairs','m2',270,''),
('R_T60S','600x600 semi-polished porcelain tiles','m2',265,''),
('R_T120','600x1200 semi-polished porcelain tiles','m2',330,''),
('R_T60N','600x600 non-slip R11 porcelain tiles','m2',250,''),
('R_TINT','Interlocking composite tiles','m2',360,''),
('R_PU','PU resin screed / coating','m2',290,''),
('R_WT','Wall tiles to washrooms','m2',240,''),
('R_EMUL','Emulsion paint (2 undercoats + finish)','m2',45,''),
('R_EXTP','Exterior paint','m2',62,''),
('R_CEILP','Emulsion to plasterboard soffits','m2',48,''),
('R_GLOSS','Gloss oil paint to woodwork','m2',72,''),
('R_PRIM','Aluminium primer to burglar-proof bars','m2',40,''),
('W','WATERPROOFING',None,None,None),
('R_WP','Liquid applied waterproofing','m2',165,''),
('R_WPS','Protective screed over waterproofing','m2',72,''),
('R_FLASH','Flashings, upstands & counter-flashings','m',185,''),
('R_GUT','Gutters / drainage channels','m',360,''),
('D','DOORS (supply & fix complete with frames/assemblies)',None,None,None)]
DRATE={'D1':18500,'D2':7800,'D3':3400,'D4':3900,'D5':16500,'D6':1900,'D7':6800,'D8':9800,'ED1':7200,'ED2':8100,'D9':6900,'MD1':14500,'FD1':4600,'M1':2900,'M2':2600,'M3':1900,'M5':2200,'GD1':12500,'DW1':15500,'DW2':33000,'SD1':14500,'SD2':8800,'SD3':7600,'SD4':9300,'SD5':7200,'RD1':23000}
for n,w,h,q in DOORS: RATES.append(('R_'+n,'Door %s %gx%g'%(n,w*1000,h*1000),'nr',DRATE[n],''))
RATES+= [('R_FRAME','50x150 hardwood door frames','m',125,''),('R_STOP','50x12 door stops','m',28,''),
('R_HINGE','Pair 100mm brass butt hinges','pr',125,''),('R_LOCK','Ordinary mortice lock','nr',480,''),('R_ILOCK','Mortice indicator lock','nr',560,''),
('G','WINDOWS (aluminium, rate = window area x base rate)',None,None,None),('R_WBASE','Aluminium glazed windows - base rate per m2','m2',1650,'Change this to re-price all windows')]
for n,w,h,q in WINS: RATES.append(('R_'+n,'Window %s %gx%g'%(n,w*1000,h*1000),'nr','=%g*%g*{R_WBASE}'%(w,h),'Linked to base rate'))
RATES+=[('L','LIFTS & PROVISIONAL SUMS',None,None,None),('R_LIFT','1000kg 12-passenger lift, stainless steel car, installed','nr',1250000,''),
('PS_JOIN','PS - joinery, wall cladding, balustrades & fittings','item',1500000,'Provisional sum'),('PS_FURN','PS - furniture & equipment','item',2500000,''),
('PS_EXT','PS - external works, infrastructure & signage','item',1800000,''),('PS_SEPT','PS - septic tank & connections','item',350000,''),('PS_BURG','PS - collapsible burglar-proofing to GF windows','item',120000,''),
('P','PRELIMINARIES (lump sums)',None,None,None)]
PREL=[(110,'P_PG','Performance guarantee',250000),(114,'P_APG','Advance payment guarantee',150000),(118,'P_INS','All risks insurance',450000),(124,'P_SET','Setting out the works',60000),
(138,'P_SIGN','Project signboard',25000),(147,'P_TRANS','Transport for workpeople',180000),(150,'P_SURV','Surveying equipment',40000),(188,'P_PLANT','Plant, tools & vehicles (crane, hoists)',1800000),
(204,'P_SUP','Contractor supervision',1200000),(217,'P_OFF','Temporary office accommodation',150000),(220,'P_FENCE','Temporary fencing & hoardings',220000),(226,'P_ROAD','Temporary roads & hardstandings',80000),
(229,'P_SAN','Temporary sanitary accommodation',60000),(232,'P_MEET','Site meetings',36000),(235,'P_AID','First aid equipment',15000),(238,'P_SEC','Site security',240000),
(242,'P_PROT','Protection of works',80000),(248,'P_WATER','Water for the works',150000),(251,'P_POWER','Temporary lighting & power',300000),(266,'P_HSE','Safety, health & welfare',120000),
(269,'P_PPE','Hard hats & protective clothing',60000),(296,'P_DEWAT','Removing water',150000),(299,'P_RUBB','Removing rubbish',120000),(304,'P_PHOTO','Progress photographs',12000),
(336,'P_TEST','Testing of materials',90000),(366,'P_OM','Operating & maintenance manuals',25000),(371,'P_HAND','Handover of completed works & cleaning',60000)]
for r,c,d,v in PREL: RATES.append((c,d,'item',v,'Lump sum'))
RS=wb.create_sheet('RATES',2)
RS['A1']='RATE LIBRARY - GHANA CEDIS (GHS)'; RS['A1'].font=Font(name=FN,size=16,bold=True)
RS['A2']='All BOQ rates link to column D below. Change a rate here and every floor bill, bill summary and the General Summary update automatically. Rates are estimates at September 2026 Accra prices, supply & fix, excl. levies/taxes.'
RS['A2'].font=Font(name=FN,size=11,italic=True); RS['A2'].alignment=Alignment(wrap_text=True); RS.merge_cells('A2:E2'); RS.row_dimensions[2].height=45
for j,h in enumerate(['CODE','DESCRIPTION','UNIT','RATE (GHS)','BASIS'],1):
    c=RS.cell(4,j,h); c.font=Font(name=FN,size=12,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(horizontal='center')
for col,w in zip('ABCDE',[12,62,8,18,40]): RS.column_dimensions[col].width=w
RROW={}; r=5
for code,desc,unit,val,basis in RATES:
    if unit is None:
        r+=1; RS.cell(r,2,desc).font=Font(name=FN,size=12,bold=True); r+=1; continue
    RROW[code]=r; r+=1
r=5
for code,desc,unit,val,basis in RATES:
    if unit is None: r+=2; continue
    RS.cell(r,1,code).font=Font(name=FN,size=11); RS.cell(r,2,desc).font=Font(name=FN,size=11); RS.cell(r,3,unit).font=Font(name=FN,size=11)
    v=val.replace('{R_WBASE}','D%d'%RROW['R_WBASE']) if isinstance(val,str) else val
    c=RS.cell(r,4,v); c.number_format='#,##0.00'; c.font=Font(name=FN,size=11,color='0000FF' if not isinstance(v,str) else '000000')
    RS.cell(r,5,basis).font=Font(name=FN,size=10,italic=True)
    r+=1
RS.freeze_panes='A5'
RREF=lambda code: '=RATES!$D$%d'%RROW[code]
# ================= item catalogue =================
CONC=('Vibrated in-situ concrete C40 as described in:',[('slab','Suspended slab 250mm thick','m3','R_CSLAB'),('slab150','Suspended roof slab 150mm thick (roof level & top of lift core)','m3','R_CSLAB'),
  ('beam','Beams (downstands below slab)','m3','R_CBEAM'),('col','Columns','m3','R_CCOL'),('wsh','Concrete wall - shear walls to main & emergency stair cores, 250mm thick','m3','R_CWALL'),
  ('wlift','Concrete wall - lift walls, 250mm thick','m3','R_CWALL'),('stair','Stairs','m3','R_CSTAIR'),('lintel','Lintels 150 x 225mm over door and window openings','m3','R_CLINT'),('coping','Copings 300 x 150mm to parapets','m3','R_CCOP')])
FORM=('E20 FORMWORK - Formwork as described to:',[('fs','Soffits of suspended slab','m2','R_FSLAB'),('fe','Edges of suspended slab 250mm high','m','R_FEDGE'),('fb','Sides and soffits of beams','m2','R_FBEAM'),
  ('fc','Columns','m2','R_FCOL'),('fw','Concrete wall (both faces)','m2','R_FWALL'),('fst','Soffits of stair flights, landings and risers','m2','R_FSTAIR'),('fl','Side and soffits of lintel / coping','m2','R_FLINT')])
REB=('E30 REINFORCEMENT - High tensile yield round bar reinforcement (standard) to BS 4449:',[('r_c10','10mm diameter in columns','t','R_Y10'),('r_c20','20mm diameter in columns','t','R_Y20'),('r_c25','25mm diameter in columns','t','R_Y25'),
  ('r_b10','10mm diameter in beams (links)','t','R_Y10'),('r_b16','16mm diameter in beams','t','R_Y16'),('r_b20','20mm diameter in beams','t','R_Y20'),('r_b25','25mm diameter in beams','t','R_Y25'),
  ('r_w10','10mm diameter in shear & lift walls (links)','t','R_Y10'),('r_w12','12mm diameter in shear, lift & ramp walls','t','R_Y12'),('r_w20','20mm diameter in shear & lift walls','t','R_Y20'),('r_w25','25mm diameter in shear & lift walls','t','R_Y25'),
  ('r_s12','12mm diameter in stairs','t','R_Y12'),('r_s16','16mm diameter in stairs','t','R_Y16'),('r_sl12','12mm diameter for suspended slab','t','R_Y12'),('r_sl16','16mm diameter for suspended slab','t','R_Y16')])
MAS=('F10 BRICK/BLOCK WALL - Solid sandcrete blockwork in cement and sand (1:4) mortar as described in:',[('blk150','150mm thick wall','m2','R_B150'),('blk200','200mm thick wall','m2','R_B200')])
WPF=('WATERPROOFING',[('wp','Liquid/applied waterproofing to roofs, terraces and wet external areas','m2','R_WP'),('wps','Protective cement-sand screed over waterproofing','m2','R_WPS'),
  ('flash','Metal/aluminium flashings, upstands and counterflashings','m','R_FLASH'),('gutter','Roof/terrace gutters and concealed drainage channels','m','R_GUT')])
DD={'D1':'1800mm x 2400mm 2-hour fire-rated double-leaf galvanised steel door (D1)','D2':'1000mm x 2200mm 1-hour fire-rated galvanised steel single-leaf security door assembly (D2)',
'D3':'900mm x 2200mm single leaf timber panel door (D3)','D4':'900mm x 2350mm single leaf timber panel door with fanlight (D4)','D5':'3200mm x 2200mm 4-panel galvanised steel framed tempered glass sliding door assembly (D5)',
'D6':'750mm x 2200mm single leaf timber flush door (D6)','D7':'700mm x 2200mm single-leaf 2-hour fire-rated galvanised steel door (D7)','D8':'1200mm x 2200mm 1-hour fire-rated leaf-and-a-half galvanised steel security door assembly (D8)',
'ED1':'1000mm x 2200mm 1-hour fire-rated galvanised steel single-leaf security door assembly - emergency stairs (ED1)','ED2':'1000mm x 2200mm 1-hour fire-rated galvanised steel single-leaf security door with translucent glass panel (ED2)',
'D9':'900mm x 2200mm 1-hour fire-rated galvanised steel single-leaf security door assembly (D9)','MD1':'1800mm x 2200mm double-leaf 2-hour fire-rated galvanised steel refuse room door (MD1)',
'FD1':'1500mm x 2200mm timber bi-fold door (FD1)','M1':'900mm x 2200mm screw-fixed galvanised steel service access panel (M1)','M2':'750mm x 2200mm screw-fixed galvanised steel service access panel (M2)',
'M3':'450mm x 2200mm screw-fixed galvanised steel service access panel (M3)','M5':'600mm x 2200mm screw-fixed galvanised steel service access panel (M5)',
'GD1':'1800mm x 2200mm double-leaf double-swing galvanised steel framed tempered glass door (GD1)','DW1':'2400mm x 2200mm double-leaf double-swing galvanised steel framed tempered glass door (DW1)',
'DW2':'5050mm x 2200mm double-leaf double-swing galvanised steel framed tempered glass door with fanlight and fixed side windows (DW2)','SD1':'3000mm x 2200mm 3-panel galvanised steel framed tempered glass sliding door (SD1)',
'SD2':'1600mm x 2200mm 2-panel galvanised steel framed tempered glass sliding door (SD2)','SD3':'1300mm x 2200mm 2-panel galvanised steel framed tempered glass sliding door (SD3)',
'SD4':'1800mm x 2200mm 2-panel galvanised steel framed tempered glass sliding door (SD4)','SD5':'1200mm x 2200mm 2-panel galvanised steel framed tempered glass sliding door (SD5)','RD1':'Metal roller shutter door 4200mm x 2157mm (RD1)'}
WDS={'W1':'4 x 2 bay alu. framed glass awning window 3600mm x 1900mm','W2':'2 bay alu. framed glass sliding window 1800mm x 1300mm','W3':'2 bay alu. framed glass sliding window 1800mm x 1500mm','W4':'1 bay alu. framed glass awning window 900mm x 1300mm',
'W5':'2 bay alu. framed glass sliding window 1800mm x 1900mm','W6':'L-shaped alu. framed fixed glass window 1750mm x 1750mm x 1500mm','W7':'2 bay alu. framed glass sliding window 1800mm x 1600mm','W8':'1 x 3 bay alu. framed glass awning window 1200mm x 1900mm',
'W9':'2 bay alu. framed glass sliding window 1200mm x 1150mm','W10':'2 bay alu. framed glass sliding window 1200mm x 1200mm','W11':'2 bay alu. framed glass sliding window 1500mm x 1300mm','W12':'2 bay alu. framed glass sliding window 1500mm x 1150mm',
'W13':'3 bay alu. framed glass awning window 2400mm x 1300mm','W14':'3 bay alu. framed glass sliding window 3600mm x 2200mm','W15':'3 bay alu. framed glass sliding window 3000mm x 1600mm','W16':'3 x 2 alu. framed glass awning window 2400mm x 1900mm',
'W17':'2 x 3 alu. framed glass awning window 1800mm x 1900mm','W18':'1 bay alu. framed glass awning window 750mm x 1300mm','W19':'3 bay alu. framed glass sliding window 2100mm x 900mm','W20':'2 x 3 alu. framed glass awning window 1450mm x 1900mm',
'HW1':'2 bay alu. framed glass sliding window 900mm x 600mm','HW2':'2 bay alu. framed glass sliding window 750mm x 600mm','HW3':'2 bay alu. framed glass sliding window 600mm x 600mm','HW4':'2 bay alu. framed glass sliding window 1200mm x 600mm'}
DOORG=('L20 DOORS - Supply and fix the following door assemblies complete with frames, ironmongery sets to steel doors and all accessories:',[('n_'+n,DD[n],'nr','R_'+n) for n,w,h,q in DOORS]+
 [('frame','50mm x 150mm frames (to timber doors)','m','R_FRAME'),('stop','50mm x 12mm door stop','m','R_STOP')])
IRON=('Ironmongery - Supply and fix the following approved ironmongery to hardwood with accompanying patented screws:',[('hinge','Pair 100mm brass butt hinges','pr','R_HINGE'),('lock','Ordinary locks','nr','R_LOCK'),('ilock','Approved high quality mortice indicator lock','nr','R_ILOCK')])
WING=("ALUMINIUM GLAZED WINDOWS - Supply and install Italian extruded aluminium glazed windows, ALCOA R50 series, powder coated aluminium profile, glazed with 6mm tinted reflective glass secured onto galvanized steel sub-frames, complete with all accessories as per Architect's detail:",
 [('n_'+n,WDS[n]+' (%s)'%n,'nr','R_'+n) for n,w,h,q in WINS])
FLR=('M10/M40 FLOOR FINISHES',[('screed','25mm cement and sand (1:3) screeded finish laid level on concrete and finished smooth with steel trowel','m2','R_SCR25'),
 ('bed','12mm cement and sand (1:4) screeded bed laid level on concrete floor to receive floor tiles','m2','R_BED12'),
 ('t_wet','400x400mm non-slip porcelain floor tiles to washrooms & WCs, adhesive, grout and cutting','m2','R_T40'),
 ('t_stair','500x500mm matte R11 porcelain tiles to staircases (treads, risers and landings)','m2','R_T50'),
 ('t_semi','600x600mm semi-polished porcelain floor tiles to habitable rooms, shops & offices','m2','R_T60S'),
 ('t_lob','600x1200mm semi-polished porcelain floor tiles to stair/lift lobbies & reception','m2','R_T120'),
 ('t_ns','600x600mm non-slip R11 porcelain tiles to corridors, terraces, verandah & service areas','m2','R_T60N'),
 ('t_int','Interlocking composite tiles to roof-top outdoor sitting & play areas','m2','R_TINT'),
 ('t_pu','Seamless polyurethane (PU) resin screed / PU coating to parking decks & ramps','m2','R_PU')])
WALLF=('M20 PLASTERED/RENDERED COATINGS - 12.5mm cement and sand (1:3) rendering as described on blockwork or concrete to:',[('rint','Interior walls','m2','R_REND'),('rwet','Interior partitions and washrooms','m2','R_REND'),('rext','External walls','m2','R_RENDX'),
 ('wtile_b','12mm cement and sand (1:4) screeded backing to receive wall tiles','m2','R_BACK12'),('wtile','Ceramic/porcelain wall tiles to washrooms to 2.4m high, adhesive & grout (type TBC)','m2','R_WT')])
PAINT=('M60 PAINTING - Prepare and apply two undercoats and finishing coat of emulsion paint on:',[('pint','Interior walls','m2','R_EMUL'),('pwet','Interior partitions and washrooms (above wall tiling)','m2','R_EMUL'),
 ('pext','Exterior walls','m2','R_EXTP'),('colpaint','Columns and walls; girth exceeding 300mm (exposed faces, ground & parking levels)','m2','R_EMUL'),
 ('ceil','Soffit of plaster board ceilings','m2','R_CEILP'),('ceilwet','Soffit of plaster board ceilings for washrooms','m2','R_CEILP'),
 ('pwood','Knot, prime, stop and apply two undercoats and one finishing coat of gloss oil paint on general surfaces of woodwork','m2','R_GLOSS'),
 ('burglar','Prepare and apply one coat of aluminium primer on mild steel burglar-proof bars (measured over window opening area)','m2','R_PRIM')])
for s in SHEETS:
    for a,b in (('wtile_b','wtile'),('pint','rint'),('pext','rext')):
        if b in Q[s]: Q[s][a]=Q[s][b]
ELEMENTS=[('ELEMENT NO. 1: IN SITU CONCRETE/LARGE PRECAST CONCRETE',[CONC,FORM,REB]),('ELEMENT NO. 2: MASONRY',[MAS]),('ELEMENT NO. 3: WATERPROOFING',[WPF]),
 ('ELEMENT NO. 4: WINDOWS/DOORS',[DOORG,IRON,WING]),('ELEMENT NO. 5: FLOOR FINISHES',[FLR]),('ELEMENT NO. 6: WALL FINISHES',[WALLF]),('ELEMENT NO. 7: PAINTING AND DECORATING',[PAINT])]
# ================= writer =================
LET='ABCDEFGHJKLMNPQRSTUVWXYZ'
def st(c,b=False,u=None,al=None,wrap=False,nf=None,border=None,sz=SZ,color=None):
    c.font=Font(name=FN,size=sz,bold=b,underline=u,color=color); c.alignment=Alignment(horizontal=al,wrap_text=wrap,vertical='top')
    if nf: c.number_format=nf
def rowbox(ws,r,top=None,bottom=None):
    for col,sd in zip('ABCDEF',[med,dbl,thin,thin,None,thin]):
        ws['%s%d'%(col,r)].border=Border(left=sd,top=top,bottom=bottom,right=med if col=='F' else None)
def newsheet(name,idx):
    if name in wb.sheetnames:
        old=wb[name]; idx=wb.sheetnames.index(name); wb.remove(old)
    ws=wb.create_sheet(name,idx)
    for col,w in zip('ABCDEFG',[13.3,120,18,12,20,26,70]): ws.column_dimensions[col].width=w
    for i in range(4):
        c=ws.cell(i+1,1,"='PRELIMS & GEN. ITEMS'!A%d"%(i+2)); st(c,al='left')
    for j,h in enumerate(['ITEM','DESCRIPTION','QTY','UNIT','RATE','AMOUNT - GHC'],1):
        c=ws.cell(5,j,h); st(c,b=True,al='center')
    c=ws.cell(5,7,'Measurement basis / calculation'); st(c,b=True,sz=12)
    for col in 'ABCDEF': ws['%s5'%col].border=Border(top=med,bottom=med,left=med if col in 'AF' else (dbl if col=='B' else thin),right=med if col=='F' else None)
    ws.page_setup.orientation='portrait'; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToHeight=0
    ws.freeze_panes='A6'
    return ws
def bill(name,idx,title,billno,elements,qsrc,notes=None):
    ws=newsheet(name,idx); r=7
    st(ws.cell(r,2,' MAIN BUILDING'),b=True,u='single'); r+=2; trow=r
    st(ws.cell(r,2,title),b=True,u='single'); r+=2
    st(ws.cell(r,2,'BILL SUMMARY'),b=True); r+=2
    # count elements with content
    used=[(et,groups) for et,groups in elements if any(k in qsrc and qsrc[k] for _,its in groups for k,_,_,_ in its)]
    sumrows={}
    for i,(et,_) in enumerate(used):
        ws.cell(r,1,LET[i]); st(ws.cell(r,1),al='center'); st(ws.cell(r,2,et)); sumrows[et]=r; rowbox(ws,r); r+=2
    tot=r; c=ws.cell(r,2,'=B%d'%trow); st(c,b=True,al='center'); st(ws.cell(r,5,'GHC'),b=True)
    c=ws.cell(r,6,'=SUM(F%d:F%d)'%(sumrows[used[0][0]] if used else r,r-1)); st(c,b=True,nf=NUM); rowbox(ws,r,top=thin,bottom=dbl)
    st(ws.cell(r+1,2,'To Bill General Summary'),al='left'); r+=4
    st(ws.cell(r,2,billno),b=True,u='single'); r+=2
    for et,groups in used:
        st(ws.cell(r,2,et),b=True,u='single'); rowbox(ws,r); r+=2; first=r; li=0
        for head,items in groups:
            its=[x for x in items if qsrc.get(x[0])]
            if not its: continue
            st(ws.cell(r,2,head),b=False,u='single',wrap=True); rowbox(ws,r); ws.row_dimensions[r].height=24*max(1,len(head)//110+1); r+=1
            for key,desc,unit,rc in its:
                terms=qsrc[key]
                ws.cell(r,1,LET[li%len(LET)]); st(ws.cell(r,1),al='center'); li+=1
                st(ws.cell(r,2,desc),wrap=True); ws.row_dimensions[r].height=24*max(1,len(desc)//105+1)
                f='='+'+'.join(terms)
                c=ws.cell(r,3,f); st(c,al='center',nf=QNUM)
                st(ws.cell(r,4,unit),al='center')
                c=ws.cell(r,5,RREF(rc)); st(c,nf=QNUM,color='000080')
                c=ws.cell(r,6,'=C%d*E%d'%(r,r)); st(c,nf=NUM)
                g=ws.cell(r,7,f[1:] if len(f)<250 else f[1:248]+'...'); g.font=Font(name=FN,size=10,color='808080'); g.alignment=Alignment(wrap_text=False)
                rowbox(ws,r); r+=2
            r+=0
        c=ws.cell(r,2,'=B%d'%(first-2)); st(c,b=True,u='single'); c=ws.cell(r,6,'=SUM(F%d:F%d)'%(first,r-1)); st(c,b=True,nf=NUM); rowbox(ws,r,top=thin,bottom=dbl)
        ws.cell(sumrows[et],6,'=F%d'%r); st(ws.cell(sumrows[et],6),nf=NUM)
        st(ws.cell(r+1,2,'Carried to Bill Summary'),al='left'); r+=4
    if notes:
        for t in notes: c=ws.cell(r,2,t); st(c,sz=12,wrap=True); c.font=Font(name=FN,size=12,italic=True,color='595959'); r+=1
    ws.print_area='A1:F%d'%r
    return ws,tot
# ================= SUBSTRUCTURE =================
SUBQ={'S1':['40.85*38.27','35.2*20.9'],'S2':['(41.0+2*2.0)*(36.0+2*2.0)'],'S3':['(45.0*40.0)*4.0'],'S4':['41.0*36.0*1.0-(69.445*0.25*1.0)-(13.294*1.0)'],'S5':['(45.0*40.0)*1.7'],
 'S6':['((45.0*40.0)-(41.0*36.0))*(4.0-1.7)'],'S7':['(45.0*40.0)*4.0-((45.0*40.0)-(41.0*36.0))*(4.0-1.7)'],'S8':['41.0*36.0*0.15'],'S9':['41.0*36.0*0.10'],'S10':['41.0*36.0*0.15'],
 'S11':['(22.05+21.825+25.57)*0.25*1.3'],'S12':['41.0*36.0*1.3'],'S20':['2*(22.05+21.825+25.57)*1.3'],'S21':['2*(41.0+36.0)'],'S22':['2*(41.0+36.0)*1.3'],'S24':['2*(41.0+36.0)*1.3']}
SUBQ['S13']=Q['SUB']['col']; SUBQ['S23']=Q['SUB']['fc']
for k in ('r_raft20','r_raft25','r_c10','r_c20','r_c25','r_w10','r_w12','r_w25'): SUBQ[k]=Q['SUB'].get(k,[])
SUBQ['S16']=SUBQ['r_c10']+SUBQ['r_w10']; SUBQ['S18']=SUBQ['r_c25']+SUBQ['r_w25']
SUBEL=[('D GROUNDWORK (D20 EXCAVATION AND FILLING)',[('Site Preparation',[('S1','Clear site of all vegetation, grub up their roots and dispose-off (Block N plot and Parking-22 forecourt)','m2','R_CLR'),
   ('S2','Excavate oversite to remove vegetable soil average 150mm deep and deposit on site for re-use average 50m from excavation','m2','R_TOP'),
   ('S3','Excavate raft foundation to reduce levels maximum depth 4.00m','m3','R_EXC'),
   ('S4','Fill to make up levels over raft average 1000mm thick with waterproof external membrane and high strength biaxial geogrid; compacted','m3','R_FILL1'),
   ('S5','Fill to make up levels below raft 1700mm thick with 300mm thick compacted crushed rock layers between geogrid layers until formation level is reached','m3','R_FILL2'),
   ('S6','Backfill and compact selected excavated material around foundations','m3','R_BACK'),('S7','Load up surplus excavated material and remove from site','m3','R_DISP'),
   ('S8','Approved imported granular material hardcore filling to make up levels: depositing and compacting in layers 150mm maximum thickness','m3','R_HARD')])]),
 ('E IN SITU CONCRETE',[('E10 Plain in-situ concrete (1:4:8 - 38mm aggregate) as described in:',[('S9','Blinding 100mm thick under raft','m3','R_BLIND'),('S10','150mm thick bed','m3','R_BED')]),
   ('Vibrated in-situ concrete C40 as described in:',[('S11','Concrete walls (raft to ground floor)','m3','R_CWALL'),('S12','Concrete in raft 1300mm thick','m3','R_CRAFT'),('S13','Concrete in columns (raft to ground floor)','m3','R_CCOL')]),
   ('E30 REINFORCEMENT - High tensile yield round bar reinforcement (standard):',[('r_raft20','20mm diameter in raft','t','R_Y20'),('r_raft25','25mm diameter in raft','t','R_Y25'),
     ('S16','10mm diameter links to columns & walls (raft to ground floor)','t','R_Y10'),('r_w12','12mm diameter to walls (raft to ground floor)','t','R_Y12'),
     ('S18','25mm diameter to columns & walls (raft to ground floor)','t','R_Y25'),('r_c20','20mm diameter in columns (raft to ground floor)','t','R_Y20')]),
   ('E20 FORMWORK',[('S20','Concrete walls (both faces)','m2','R_FWALL'),('S21','Edges of bed 150mm high','m','R_FEDGE'),('S22','Edges of raft 1300mm high','m2','R_FRAFT'),('S23','Sides of columns (raft to ground floor)','m2','R_FCOL')])]),
 ('BLOCK WORK',[('Solid sandcrete blockwork in cement and sand (1:4) mortar:',[('S24','150mm walls to plinth','m2','R_B150')])])]
idx=wb.sheetnames.index('SUBSTRUCTURE')
subws,subtot=bill('SUBSTRUCTURE',idx,'SUBSTRUCTURE (ALL PROVISIONAL)','BILL NR. 2 - SUBSTRUCTURE',SUBEL,SUBQ,
   notes=['Raft top SSL -1.000; formation -4.000 (STR-001). Raft 41.0 x 36.0m measured on STR-001/STR-002.'])
# ================= FLOORS =================
FL=[('GROUND FLR ','GF',' GROUND FLOOR (INCL. M1 PARKING LEVEL 1B)','BILL NR. 3 - GROUND FLOOR'),('1ST FLR ','1ST',' FIRST FLOOR - LEVEL 2A PARKING (INCL. M2 LEVEL 2B)','BILL NR. 4 - FIRST FLOOR'),
 ('2ND FLR','2ND',' SECOND FLOOR - LEVEL 3A PARKING (INCL. MEZZANINE LEVEL 3B)','BILL NR. 5 - SECOND FLOOR'),('3RD FLR','3RD',' THIRD FLOOR','BILL NR. 6 - THIRD FLOOR'),
 ('4TH FLR ','4TH',' FOURTH FLOOR','BILL NR. 7 - FOURTH FLOOR'),('5TH FLR ','5TH',' FIFTH FLOOR','BILL NR. 8 - FIFTH FLOOR'),('6TH FLR  ','6TH',' SIXTH FLOOR','BILL NR. 9 - SIXTH FLOOR'),
 ('7TH FLR   ','7TH',' SEVENTH FLOOR','BILL NR. 10 - SEVENTH FLOOR'),('8TH FLR   ','8TH',' EIGHTH FLOOR','BILL NR. 11 - EIGHTH FLOOR'),('9TH FLR    ','9TH',' NINTH FLOOR','BILL NR. 12 - NINTH FLOOR'),
 ('10TH FLR   ','10TH',' TENTH FLOOR','BILL NR. 13 - TENTH FLOOR'),('11TH FLR','11TH',' ELEVENTH LEVEL - ROOF TOP TERRACE','BILL NR. 14 - ROOF TOP TERRACE'),('12TH FLR ','12TH',' TWELFTH LEVEL - ROOF LEVEL & TOP OF LIFT CORE','BILL NR. 15 - ROOF LEVEL')]
names=wb.sheetnames
TOT={}
NOTE=['Doors and windows: schedules ARCH-501/502 give building totals only (no tags on plans); quantities are allocated to floors pro-rata to measured habitable / washroom areas. Building totals equal the schedules.']
for nm,code,title,billno in FL:
    real=[n for n in names if n.strip()==nm.strip()][0]
    ws,t=bill(real,wb.sheetnames.index(real),title,billno,ELEMENTS,Q[code],notes=NOTE); TOT[code]=(real,t)
# ================= LIFTS & PS =================
LQ={'lift':['2'],'pj':['1'],'pf':['1'],'pe':['1'],'pst':['1'],'pb':['1']}
LEL=[('LIFTS',[('Lift installation',[('lift','1000kg stainless steel car finish passenger lift for 12 passengers, serving all floors, complete (LIFT 1 & LIFT 2)','nr','R_LIFT')])]),
 ('PROVISIONAL SUMS',[('Provide the following provisional sums:',[('pj','Provisional sum for joinery, wall cladding, balustrade and other fitting works','item','PS_JOIN'),
   ('pf','Provisional sum for furniture and equipment (wardrobes, kitchen cabinets, reception desk and chairs)','item','PS_FURN'),
   ('pe','Provisional sum for external works / external infrastructure / signage (Axonopus compressus, kerbs, boundary structures, aprons, street lights, flood lights)','item','PS_EXT'),
   ('pst','Provisional sum for septic tank and other connections','item','PS_SEPT'),('pb','Allow a sum for metal collapsible burglar-proofing to ground floor windows including painting','item','PS_BURG')])])]
lws,ltot=bill('LIFTS & PS',len(wb.sheetnames),' LIFTS AND PROVISIONAL SUMS','BILL NR. 16 - LIFTS AND PROVISIONAL SUMS',LEL,LQ,
   notes=['Service/goods lift in the lump BOQ is not shown on the drawings (only two lift shafts) - omitted pending confirmation.'])
# ================= PRELIMS =================
P=wb['PRELIMS & GEN. ITEMS']
for r,c,d,v in PREL:
    cc=P.cell(r,6,'=RATES!$D$%d'%RROW[c]); cc.number_format=NUM; cc.font=Font(name=FN,size=P.cell(r,2).font.sz or 12)
pages=[(43,74),(75,106),(107,137),(138,163),(164,186),(187,216),(217,241),(242,268),(269,295),(296,325),(326,355),(356,370),(371,379)]
for a,b in pages:
    P.cell(b,6,'=SUM(F%d:F%d)'%(a,b-1)).number_format=NUM
P['G41']=None
# ================= GEN SUMMARY =================
G=wb['GEN SUMMARY']
G['D9']='=SUBSTRUCTURE!F%d'%subtot
labels={'GF':(11,'GROUND FLOOR (INCL. M1 PARKING LEVEL)'),'1ST':(13,'FIRST FLOOR (INCL. M2 PARKING LEVEL)'),'2ND':(15,'SECOND FLOOR (INCL. MEZZANINE PARKING LEVEL)'),
 '3RD':(17,'THIRD FLOOR'),'4TH':(19,'FOURTH FLOOR'),'5TH':(21,'FIFTH FLOOR'),'6TH':(23,'SIXTH FLOOR'),'7TH':(25,'SEVENTH FLOOR'),'8TH':(27,'EIGHTH FLOOR'),
 '9TH':(29,'NINTH FLOOR'),'10TH':(31,'TENTH FLOOR'),'11TH':(33,'ELEVENTH LEVEL - ROOF TOP TERRACE'),'12TH':(35,'TWELFTH LEVEL - ROOF LEVEL & LIFT CORE')}
fmtD=G['D7'].number_format
for code,(row,lab) in labels.items():
    real,t=TOT[code]; G.cell(row,2,lab); c=G.cell(row,4,"='%s'!F%d"%(real,t)); c.number_format=fmtD; c.font=copy.copy(G['D7'].font)
G['A25']='K'; G['A33']='P'; G['A35']='Q'
G['A37']='R'; G['B37']='LIFTS AND PROVISIONAL SUMS'; G['B37'].font=copy.copy(G['B35'].font); G['A37'].font=copy.copy(G['A35'].font)
c=G['D37']; c.value="='LIFTS & PS'!F%d"%ltot; c.number_format=fmtD; c.font=copy.copy(G['D7'].font)
G['D38']='=SUM(D6:D37)'
G['B41']='ADD: STATUTORY LEVIES AND TAXES'; G['D41']='=D38*C41'; G['D41'].number_format=fmtD
G['D2']=None
wb.move_sheet('RATES',offset=len(wb.sheetnames)-1-wb.sheetnames.index('RATES'))
# strip stale defined names / external links inherited from the template
import re as _re
for ws_ in wb.worksheets:
    for row in ws_.iter_rows():
        for c in row:
            if isinstance(c.value,str) and _re.search(r"\[\d+\]",c.value): print('EXT REF',ws_.title,c.coordinate,c.value[:60])
for n in list(wb.defined_names): del wb.defined_names[n]
for ws_ in wb.worksheets:
    for n in list(ws_.defined_names): del ws_.defined_names[n]
wb._external_links=[]
wb.calculation.fullCalcOnLoad=True
out='/home/user/OnukpaTay/BOQ_KANESHIE_BLOCK_N_PRICED_FLOOR_BY_FLOOR.xlsx'
wb.save(out); print('saved',out, wb.sheetnames)
