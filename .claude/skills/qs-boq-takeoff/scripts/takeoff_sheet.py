#!/usr/bin/env python3
"""Add a TAKEOFF (dimension sheet) to an existing BOQ workbook and link every QTY cell to it.

The client's workbook is edited at XML level (not re-saved with openpyxl), so every existing
sheet, style, image, print setting and add-in stays byte-identical except the QTY formulas.

Each QTY formula, e.g. =5*0.75*0.5*3.05+2*(41+36)*1.3-13.294, is split into dimension rows
(Times | Dim1 | Dim2 | Dim3 -> PRODUCT), collections like (10*(1+0.3)+19*(0.9+0.3))*0.15*0.225
are distributed into one row per term, deductions become negative rows, and each item ends in a
"Total carried to BOQ" SUM row that the bill's QTY cell references.

usage:
  takeoff_sheet.py IN.xlsx OUT.xlsx [--sheets "SUBSTRUCTURE,GROUND FLR "] [--labels labels.json]
                   [--item-col A --desc-col B --qty-col C --unit-col D] [--name TAKEOFF]

labels.json (optional) gives meaningful dimension descriptions:
  {"<sheet>!<cell>": [{"expr": "5*0.75*0.5*3.05", "label": "COL-01A x5, 3rd to 4th"}, ...]}
  When given, the item is built from these expressions (they must sum to the QTY value).
Without labels, every row is labelled from its expression.
"""
import argparse, ast, collections, math, os, re, shutil, tempfile, zipfile
from xml.sax.saxutils import escape, quoteattr
import openpyxl
from openpyxl.utils import column_index_from_string as ci

ADDS = (ast.Add, ast.Sub)

# ---------------------------------------------------------------- expression handling
def parse(e):
    e = e.lstrip('=').strip()
    t = ast.parse(e, mode='eval').body
    toks = iter(re.findall(r'\d+\.?\d*(?:[eE][-+]?\d+)?', e))
    for x in sorted([x for x in ast.walk(t) if isinstance(x, ast.Constant)], key=lambda x: (x.lineno, x.col_offset)):
        x._s = next(toks)                         # keep the original spelling of numbers
    return t

def txt(n):
    if isinstance(n, ast.Constant):
        return getattr(n, '_s', None) or repr(n.value)
    if isinstance(n, ast.BinOp):
        l, r = txt(n.left), txt(n.right)
        wrap = lambda x, s: '(' + s + ')' if isinstance(x, ast.BinOp) and isinstance(x.op, ADDS) else s
        if isinstance(n.op, ADDS):
            return l + ('+' if isinstance(n.op, ast.Add) else '-') + wrap(n.right, r)
        if isinstance(n.op, ast.Mult): return wrap(n.left, l) + '*' + wrap(n.right, r)
        if isinstance(n.op, ast.Div): return wrap(n.left, l) + '/' + ('(' + r + ')' if isinstance(n.right, ast.BinOp) else r)
    if isinstance(n, ast.UnaryOp): return '-' + txt(n.operand)
    return ast.unparse(n)

def jf(fs): return '*'.join('(' + txt(f) + ')' if isinstance(f, ast.BinOp) and isinstance(f.op, ADDS) else txt(f) for f in fs)

def terms_of(n, s=1):
    if isinstance(n, ast.BinOp) and isinstance(n.op, ADDS):
        return terms_of(n.left, s) + terms_of(n.right, s if isinstance(n.op, ast.Add) else -s)
    if isinstance(n, ast.UnaryOp) and isinstance(n.op, ast.USub): return terms_of(n.operand, -s)
    return [(s, n)]

def facs(n):
    if isinstance(n, ast.BinOp) and isinstance(n.op, ast.Mult): return facs(n.left) + facs(n.right)
    return [n]

def is_collection(n):
    """A bracketed sum that is a list of measured items (distribute it), as opposed to a girth like (0.75+0.5)."""
    if not (isinstance(n, ast.BinOp) and isinstance(n.op, ADDS)): return False
    T = terms_of(n)
    if len(T) >= 3: return True
    for s, t in T:
        f = facs(t)
        if len(f) >= 3 or any(isinstance(x, ast.BinOp) and isinstance(x.op, ADDS) for x in f) \
           or (isinstance(t, ast.BinOp) and isinstance(t.op, ast.Div)):
            return True
    return False

def expand(n, top=False):
    """-> list of (sign, [factor nodes]) dimension rows"""
    if isinstance(n, ast.BinOp) and isinstance(n.op, ADDS) and (top or is_collection(n)):
        return [(s * s2, f) for s, t in terms_of(n) for s2, f in expand(t)]
    if isinstance(n, ast.UnaryOp) and isinstance(n.op, ast.USub):
        return [(-s, f) for s, f in expand(n.operand)]
    if isinstance(n, ast.BinOp) and isinstance(n.op, ast.Mult):
        return [(a * b, f + g) for a, f in expand(n.left) for b, g in expand(n.right)]
    if isinstance(n, ast.BinOp) and isinstance(n.op, ast.Div):
        return [(s, f[:-1] + [ast.BinOp(left=f[-1], op=ast.Div(), right=n.right)]) for s, f in expand(n.left)]
    return [(1, [n])]

isint = lambda n: isinstance(n, ast.Constant) and isinstance(n.value, int)

def layout(fs):
    """factors -> (times text or None, [dim texts]) ; a leading integer is 'times'."""
    fs = list(fs); nr = []
    if len(fs) >= 2 and isint(fs[0]):
        nr.append(fs.pop(0))
        if len(fs) == 3 and isint(fs[0]) and all(isinstance(x, ast.Constant) for x in fs[1:]): nr.append(fs.pop(0))
    dims, i = [], 0
    while i < len(fs):                            # 2*(a+b) stays together as one girth
        if isint(fs[i]) and fs[i].value == 2 and i + 1 < len(fs) and isinstance(fs[i + 1], ast.BinOp) and isinstance(fs[i + 1].op, ADDS):
            dims.append(ast.BinOp(left=fs[i], op=ast.Mult(), right=fs[i + 1])); i += 2
        else:
            dims.append(fs[i]); i += 1
    while len(dims) > 3:
        dims = [ast.BinOp(left=dims[0], op=ast.Mult(), right=dims[1])] + dims[2:]
    return ('*'.join(txt(x) for x in nr) if nr else None), [txt(d) for d in dims]

def ev(s): return eval(s, {'__builtins__': {}})
isnum = lambda s: re.fullmatch(r'-?\d+(\.\d+)?', s) is not None

# ---------------------------------------------------------------- build row model
def build(src, sheets, cols, labels):
    wb = openpyxl.load_workbook(src)
    wbv = openpyxl.load_workbook(src, data_only=True)
    A, B, C, D = (ci(cols[k]) for k in ('item', 'desc', 'qty', 'unit'))
    R, links, bills, checks = [], {}, [], []
    def row(kind, **c): R.append(dict(kind=kind, c=c)); return len(R)
    for nm in sheets:
        ws = wb[nm]
        title = next((ws.cell(r, B).value.strip() for r in range(1, 60)
                      if isinstance(ws.cell(r, B).value, str) and ws.cell(r, B).value.strip().upper().startswith('BILL NR')), nm.strip())
        bstart = row('bill', A=title); bills.append((nm, title, bstart))
        hdr = next((r for r in range(1, 40) if str(ws.cell(r, C).value).strip().upper() in ('QTY', 'QUANTITY')), 1)
        pending = []
        for r in range(hdr + 1, ws.max_row + 1):
            a, b, v, u = (ws.cell(r, k).value for k in (A, B, C, D))
            if isinstance(v, str) and v.startswith('='):
                if re.search(r'[A-Za-z!$]', v.lstrip('=')):      # already a reference / function: leave alone
                    continue
                for h in pending: row('head', C=h)
                pending = []
                d = b.strip() if isinstance(b, str) else ''
                row('item', A=nm.strip(), B=str(a or ''), C=d, I=str(u or ''), K=("'%s'!%s%d" % (nm, cols['qty'], r)))
                first = len(R) + 1
                key = '%s!%s%d' % (nm, cols['qty'], r)
                if key in labels:
                    src_terms = [(1, x['expr'], x.get('label', '')) for x in labels[key]]
                else:
                    src_terms = [(s, txt(t), None) for s, t in terms_of(parse(v))]
                for s, e, lb in src_terms:
                    rows = [(s * s2, f) for s2, f in expand(parse(e), top=True)]
                    for k, (s2, fs) in enumerate(rows):
                        lab = lb if lb and len(rows) == 1 else ('%s (%d)' % (lb, k + 1) if lb else
                              'Dimension: ' + jf(fs).replace('*', ' x '))
                        if s2 < 0 and not lab.lower().startswith('ddt'): lab = 'Ddt  ' + lab
                        cells = dict(C=lab, J=('-' if s2 < 0 else '') + jf(fs))
                        nrt, dims = layout(fs)
                        if len(dims) == 1 and not nrt:
                            cells['H'] = (('-' if s2 < 0 else '') + dims[0], 'const')
                        else:
                            if nrt: cells['D'] = nrt
                            for col, dd in zip('EFG', dims): cells[col] = dd
                            cells['H'] = ('-' if s2 < 0 else '', 'prod')
                        row('dim', **cells)
                t = row('total', C='Total carried to BOQ  %s item %s' % (nm.strip(), a or ''), H=(first, len(R)), I=str(u or ''))
                links[(nm, r)] = t
                checks.append((nm, r, wbv[nm].cell(r, C).value, t))
            elif isinstance(b, str) and not a and not b.startswith('=') and len(b) < 400:
                if re.match(r'\s*(Carried|To Bill|CF|BF|BILL SUMMARY|BILL NR)', b, re.I): pending = []
                else: pending.append(b.strip())
        row('blank')
    return R, links, bills, checks

# ---------------------------------------------------------------- styles (appended, never edited)
def add_styles(st):
    def addlist(tag, items):
        nonlocal st
        m = re.search(r'<%s count="(\d+)"[^>]*>' % tag, st); n = int(m.group(1))
        end = st.index('</%s>' % tag, m.end())
        st = st[:end] + ''.join(items) + st[end:]
        st = st[:m.start()] + m.group(0).replace('count="%d"' % n, 'count="%d"' % (n + len(items))) + st[m.end():]
        return list(range(n, n + len(items)))
    def font(sz=11, b=False, i=False, u=False, color=None):
        return '<font>' + ('<b/>' if b else '') + ('<i/>' if i else '') + ('<u/>' if u else '') + '<sz val="%g"/>' % sz + \
               ('<color rgb="%s"/>' % color if color else '<color theme="1"/>') + '<name val="Calibri"/><family val="2"/><scheme val="minor"/></font>'
    FN = addlist('fonts', [font(), font(b=True), font(14, b=True, color='FF1F3864'), font(10, i=True, color='FF595959'), font(b=True, color='FFFFFFFF'),
                           font(u=True, color='FF0563C1'), font(9, color='FF7F7F7F'), font(12, b=True, color='FF1F3864'), font(i=True, b=True, color='FF1F3864')])
    fill = lambda c: '<fill><patternFill patternType="solid"><fgColor rgb="%s"/><bgColor indexed="64"/></patternFill></fill>' % c
    PL = addlist('fills', [fill('FF1F3864'), fill('FFD9E1F2'), fill('FFFFF2CC'), fill('FFF2F2F2')])
    thin = lambda c: ''.join('<%s style="thin"><color rgb="%s"/></%s>' % (s, c, s) for s in ('left', 'right', 'top', 'bottom')) + '<diagonal/>'
    BD = addlist('borders', ['<border>%s</border>' % thin('FF808080'),
                             '<border><left/><right/><top style="thin"><color auto="1"/></top><bottom style="double"><color auto="1"/></bottom><diagonal/></border>',
                             '<border><left/><right/><top/><bottom style="medium"><color rgb="FF1F3864"/></bottom><diagonal/></border>',
                             '<border>%s</border>' % thin('FFBFBFBF')])
    if '<numFmts' in st:
        m = re.search(r'<numFmts count="(\d+)">', st)
        st = st.replace(m.group(0), '<numFmts count="%d">' % (int(m.group(1)) + 2))
        st = st.replace('</numFmts>', '<numFmt numFmtId="190" formatCode="#,##0.000"/><numFmt numFmtId="191" formatCode="#,##0.###"/></numFmts>')
    else:
        st = re.sub(r'(<styleSheet[^>]*>)', r'\1<numFmts count="2"><numFmt numFmtId="190" formatCode="#,##0.000"/><numFmt numFmtId="191" formatCode="#,##0.###"/></numFmts>', st, 1)
    def xf(f, fl=0, b=0, nf=0, h=None, v='center', wrap=False, ind=0):
        al = '<alignment' + (' horizontal="%s"' % h if h else '') + (' vertical="%s"' % v if v else '') + (' wrapText="1"' if wrap else '') + (' indent="%d"' % ind if ind else '') + '/>'
        return '<xf numFmtId="%d" fontId="%d" fillId="%d" borderId="%d" xfId="0" applyNumberFormat="1" applyFont="1" applyFill="1" applyBorder="1" applyAlignment="1">%s</xf>' % (nf, f, fl, b, al)
    names = ['title', 'note', 'hdr', 'bill', 'head', 'itemA', 'itemD', 'unit', 'text', 'in', 'nr', 'qty', 'totl', 'totn', 'totu', 'calc', 'link', 'linkb']
    xs = [xf(FN[2]), xf(FN[3]), xf(FN[4], PL[0], BD[0], h='center', wrap=True), xf(FN[7], PL[1], BD[2]), xf(FN[8], ind=1),
          xf(FN[1], h='center', v='top'), xf(FN[1], wrap=True, v='top'), xf(FN[0], h='center'), xf(FN[0], ind=2), xf(FN[0], PL[2], BD[3], nf=190),
          xf(FN[0], PL[2], BD[3], nf=191, h='center'), xf(FN[0], nf=190), xf(FN[1], PL[3], BD[1], ind=1), xf(FN[1], PL[3], BD[1], nf=190),
          xf(FN[1], PL[3], BD[1], h='center'), xf(FN[6]), xf(FN[5]), xf(FN[5], v='top')]
    ids = addlist('cellXfs', xs)        # must run before st is returned
    return st, dict(zip(names, ids))

# ---------------------------------------------------------------- write
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('out')
    ap.add_argument('--sheets', help='comma list; default = every sheet that has a QTY header and numeric QTY formulas')
    ap.add_argument('--labels'); ap.add_argument('--name', default='TAKEOFF')
    for k, d in (('item', 'A'), ('desc', 'B'), ('qty', 'C'), ('unit', 'D')): ap.add_argument('--%s-col' % k, default=d)
    a = ap.parse_args()
    cols = dict(item=a.item_col, desc=a.desc_col, qty=a.qty_col, unit=a.unit_col)
    labels = __import__('json').load(open(a.labels)) if a.labels else {}
    wb = openpyxl.load_workbook(a.src)
    if a.sheets:
        sheets = [s for s in wb.sheetnames if s.strip() in [x.strip() for x in a.sheets.split(',')]]
    else:
        q = ci(cols['qty'])
        sheets = [ws.title for ws in wb if any(str(ws.cell(r, q).value).strip().upper() == 'QTY' for r in range(1, 40))
                  and any(isinstance(ws.cell(r, q).value, str) and ws.cell(r, q).value.startswith('=') and not re.search(r'[A-Za-z!$]', ws.cell(r, q).value[1:])
                          for r in range(1, ws.max_row + 1))]
    R, links, bills, checks = build(a.src, sheets, cols, labels)

    W = tempfile.mkdtemp(); zipfile.ZipFile(a.src).extractall(W)
    rd = lambda p: open(os.path.join(W, p), encoding='utf-8', newline='').read()
    wr = lambda p, s: open(os.path.join(W, p), 'w', encoding='utf-8', newline='').write(s)
    st, S = add_styles(rd('xl/styles.xml')); wr('xl/styles.xml', st)

    HDR = 4; IDX0 = HDR + 1; OFF = IDX0 + len(bills) + 2 - 1
    rowof = lambda i: i + OFF
    rows, HL, VAL = [], [], {}
    cs = lambda c, r, t, s: '<c r="%s%d" s="%d" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (c, r, s, escape(t))
    cn = lambda c, r, v, s: '<c r="%s%d" s="%d"><v>%r</v></c>' % (c, r, s, v)
    cf = lambda c, r, f, v, s: '<c r="%s%d" s="%d"><f>%s</f><v>%r</v></c>' % (c, r, s, escape(f), v)
    rows.append((2, [cs('A', 2, 'QUANTITY TAKE-OFF  -  DIMENSION SHEETS', S['title'])], 22))
    rows.append((3, [cs('A', 3, 'Yellow cells are dimensions - edit them and the QTY column of each bill (and all amounts / summaries) updates. '
                     'Qty = Times x Dim1 x Dim2 x Dim3 (blank = 1); Ddt lines are negative.', S['note'])], None))
    H = ['BILL / FLOOR', 'ITEM', 'DESCRIPTION / LOCATION OF DIMENSION', 'TIMES / NR', 'LENGTH / DIM 1', 'WIDTH / DIM 2', 'HEIGHT / DEPTH / DIM 3',
         'QUANTITY', 'UNIT', 'ORIGINAL CALCULATION (FOR REFERENCE)', 'LINK TO BOQ QTY CELL']
    rows.append((HDR, [cs(c, HDR, h, S['hdr']) for c, h in zip('ABCDEFGHIJK', H)], 32))
    rows.append((IDX0, [cs('A', IDX0, 'BILL INDEX (click to jump)', S['itemD'])], None))
    for k, (nm, title, bstart) in enumerate(bills):
        r = IDX0 + 1 + k
        rows.append((r, [cs('A', r, '   ' + title, S['link']), cs('K', r, "Open bill: '%s'" % nm.strip(), S['link'])], None))
        HL += [('A%d' % r, "%s!A%d" % (a.name, rowof(bstart))), ('K%d' % r, "'%s'!A1" % nm)]
    for i, x in enumerate(R, 1):
        r, c, k, out, ht = rowof(i), x['c'], x['kind'], [], None
        if k == 'bill':
            out = [cs('A', r, c['A'], S['bill'])] + [cs(col, r, '', S['bill']) for col in 'BCDEFGHIJK']; ht = 21
        elif k == 'head':
            out = [cs('C', r, c['C'], S['head'])]
        elif k == 'item':
            ln = max(1, math.ceil(len(c['C']) / 78)); ht = 15.6 * ln + 1 if ln > 1 else None
            out = [cs('A', r, c['A'], S['itemA']), cs('B', r, c['B'], S['itemA']), cs('C', r, c['C'], S['itemD']), cs('I', r, c['I'], S['unit']), cs('K', r, c['K'], S['linkb'])]
            HL.append(('K%d' % r, c['K']))
        elif k == 'dim':
            out = [cs('C', r, c['C'], S['text'])]; prod = 1.0
            for col in 'DEFG':
                if col in c:
                    t = c[col]; v = ev(t); prod *= v; s = S['nr' if col == 'D' else 'in']
                    out.append(cn(col, r, v, s) if isnum(t) else cf(col, r, t, v, s))
            h, typ = c['H']
            if typ == 'prod':
                v = -prod if h == '-' else prod
                out.append(cf('H', r, ('-' if h == '-' else '') + 'PRODUCT(D%d:G%d)' % (r, r), v, S['qty']))
            else:
                v = ev(h); out.append(cn('H', r, v, S['in']) if isnum(h) else cf('H', r, h, v, S['in']))
            VAL[i] = v; out.append(cs('J', r, c['J'], S['calc']))
        elif k == 'total':
            f0, f1 = c['H']; v = sum(VAL.get(j, 0) for j in range(f0, f1 + 1)); VAL[i] = v
            out = [cs('C', r, c['C'], S['totl'])] + [cs(col, r, '', S['totl']) for col in 'DEFG'] + \
                  [cf('H', r, 'SUM(H%d:H%d)' % (rowof(f0), rowof(f1)), v, S['totn']), cs('I', r, c['I'], S['totu'])]
        rows.append((r, out, ht))
    bad = [(nm, r, o, VAL[t]) for nm, r, o, t in checks if o is None or abs(o - VAL[t]) > 1e-6 * max(1, abs(o))]
    for b in bad: print('MISMATCH (sheet,row,original,takeoff):', b)
    if bad: raise SystemExit('take-off totals do not reproduce the original quantities - nothing written')

    last = max(r for r, _, _ in rows)
    colw = dict(A=13, B=6, C=78, D=10, E=12, F=12, G=13, H=15, I=7, J=48, K=24)
    sx = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
          '<sheetPr><tabColor rgb="FF1F3864"/><pageSetUpPr fitToPage="1"/></sheetPr><dimension ref="A1:K%d"/>' % last,
          '<sheetViews><sheetView workbookViewId="0" zoomScale="85" zoomScaleNormal="85"><pane ySplit="%d" topLeftCell="A%d" activePane="bottomLeft" state="frozen"/>'
          '<selection pane="bottomLeft" activeCell="A%d" sqref="A%d"/></sheetView></sheetViews>' % (HDR, HDR + 1, HDR + 1, HDR + 1),
          '<sheetFormatPr defaultRowHeight="15"/><cols>' + ''.join('<col min="%d" max="%d" width="%g" customWidth="1"/>' % (i, i, colw[c]) for i, c in enumerate('ABCDEFGHIJK', 1)) + '</cols><sheetData>']
    for r, cells, ht in sorted(rows):
        sx.append('<row r="%d"%s>' % (r, ' ht="%g" customHeight="1"' % ht if ht else '') + ''.join(cells) + '</row>')
    sx.append('</sheetData><hyperlinks>' + ''.join('<hyperlink ref="%s" location=%s display=%s/>' % (ref, quoteattr(loc), quoteattr(loc)) for ref, loc in HL) + '</hyperlinks>')
    sx.append('<pageMargins left="0.4" right="0.4" top="0.6" bottom="0.6" header="0.3" footer="0.3"/><pageSetup paperSize="9" orientation="landscape" fitToHeight="0"/></worksheet>')

    wbx, rels, ct = rd('xl/workbook.xml'), rd('xl/_rels/workbook.xml.rels'), rd('[Content_Types].xml')
    sheetfile = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]*)"[^>]*r:id="(rId\d+)"', wbx):
        tgt = re.search(r'Id="%s"[^>]*Target="([^"]*)"' % m.group(2), rels) or re.search(r'Target="([^"]*)"[^>]*Id="%s"' % m.group(2), rels)
        sheetfile[m.group(1).replace('&amp;', '&')] = tgt.group(1).lstrip('/').replace('xl/', '', 1)
    n_new = 1 + max(int(x) for x in re.findall(r'worksheets/sheet(\d+)\.xml', rels))
    new_part = 'xl/worksheets/sheet%d.xml' % n_new
    wr(new_part, ''.join(sx))
    by = collections.defaultdict(list)
    for (nm, r), t in links.items(): by[nm].append((r, t))
    for nm, lst in by.items():
        p = 'xl/' + sheetfile[nm]; x = rd(p)
        for r, t in lst:
            x, k = re.subn(r'(<c r="%s%d"[^>]*>)<f>[^<]*</f>' % (cols['qty'], r), lambda m: m.group(1) + '<f>%s!$H$%d</f>' % (a.name, rowof(t)), x)
            assert k == 1, (nm, r)
        wr(p, x)
    rid = 'rId%d' % (1 + max(int(x) for x in re.findall(r'Id="rId(\d+)"', rels)))
    sid = 1 + max(int(x) for x in re.findall(r'sheetId="(\d+)"', wbx))
    wbx = wbx.replace('</sheets>', '<sheet name="%s" sheetId="%d" r:id="%s"/></sheets>' % (a.name, sid, rid))
    wbx = re.sub(r'<calcPr([^>]*?)(/?)>', lambda m: '<calcPr%s fullCalcOnLoad="1"%s>' % (m.group(1).replace(' fullCalcOnLoad="1"', ''), m.group(2)), wbx, 1)
    nsheets = len(re.findall(r'<sheet ', wbx)) - 1
    pt = '<definedName name="_xlnm.Print_Titles" localSheetId="%d">%s!$%d:$%d</definedName>' % (nsheets, a.name, HDR, HDR)
    wbx = wbx.replace('</definedNames>', pt + '</definedNames>') if '</definedNames>' in wbx else wbx.replace('</sheets>', '</sheets><definedNames>%s</definedNames>' % pt, 1)
    wr('xl/workbook.xml', wbx)
    rels = re.sub(r'<Relationship [^>]*calcChain[^>]*/>', '', rels)
    rels = rels.replace('</Relationships>', '<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet%d.xml"/></Relationships>' % (rid, n_new))
    wr('xl/_rels/workbook.xml.rels', rels)
    ct = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', '', ct)
    ct = ct.replace('</Types>', '<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>' % new_part)
    wr('[Content_Types].xml', ct)
    names = [i.filename for i in zipfile.ZipFile(a.src).infolist() if i.filename != 'xl/calcChain.xml'] + [new_part]
    with zipfile.ZipFile(a.out, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in names: z.write(os.path.join(W, f), f)
    shutil.rmtree(W)
    print('sheets:', len(sheets), '| items linked:', len(links), '| take-off rows:', last, '| written', a.out)

if __name__ == '__main__':
    main()
