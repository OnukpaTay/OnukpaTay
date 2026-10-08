"""Priority Insurance - Bill No. 1 Preliminaries & General Items, built on the Thoroughbred Place BOQ."""
import sys, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ref, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(ref, keep_links=False)
areas = {ws.title: ws.print_area for ws in wb}
for n in list(wb.defined_names.keys()): del wb.defined_names[n]
for ws in wb:
    for n in list(ws.defined_names.keys()): del ws.defined_names[n]
for ws in wb:
    if areas[ws.title]: ws.print_area = areas[ws.title].split('!')[-1].replace('$', '')

HELP_HDR = Font(name='Arial Narrow', bold=True, color='1F3864')
INFILL = PatternFill('solid', fgColor='FFF2CC')

# ---------------- Cover ----------------
cp = wb['Cover Page ']
cp['B11'] = 'PROPOSED \nBILLS OF QUANTITIES FOR PRIORITY INSURANCE, PROPOSED HEAD OFFICE - RIDGE, ACCRA'
cp['F23'] = 'DATE:08/10/2026'

# ---------------- GEN SUMMARY ----------------
gs = wb['GEN SUMMARY']
for row in gs.iter_rows(min_row=1, max_row=70, max_col=12):
    for c in row:
        if c.row >= 6 or c.column > 1: c.value = None
gs['A1'] = 'PROJECT: PROPOSED HEAD OFFICE FOR PRIORITY INSURANCE'
gs['A2'] = 'LOCATION: RIDGE, ACCRA'
gs['A3'] = 'CLIENT: PRIORITY INSURANCE'
gs['A4'] = "='Cover Page '!F23"
gs['A5'] = 'ITEM'; gs['B5'] = 'GENERAL SUMMARY'; gs['D5'] = 'AMOUNT USD'
items = [  # Preliminary Estimate rev 6 (OPTION 1rev col E), before discount
 ('BILL NO. 1- PRELIMINARIES AND GENERAL ITEMS', '=366625+50000'),
 ('BILL N0. 2- DEMOLITIONS AND ALTERATIONS', 50000),
 ('BILL NO. 3 -SUBSTRUCTURE (ALL PROVISIONAL)', '=1585*350'),
 ('BILL NO. 4 -SUB-BASEMENT FLOOR ', '=1585*400'),
 ('BILL NO. 5 - GROUND FLOOR', '=450*500'),
 ('BILL NO. 6 - FIRST FLOOR ', '=600*500'),
 ('BILL NO. 7- SECOND FLOOR', '=600*500'),
 ('BILL NO. 8 -THIRD FLOOR', '=600*500'),
 ('BILL NO. 9 -FOURTH FLOOR', '=600*500'),
 ('BILL NO. 10 - FIFTH FLOOR ', '=600*500'),
 ('BILL NO. 11 - KITCHEN CABINETRY/WALL CLADDING', 134515.22004867764),
 ('BILL NO. 12. PLUMBING, FIRE-FIGHTING AND MECHANICAL ', '=659925+125000'),
 ('BILL NO.13: ELECTRICAL WORKS', '=1026550+50000'),
 ('BILL NO. 14. EXTERNAL FAÇADE WORKS', '=606968.250712613+50000'),
 ('BILL NO. 15.- EXTERNAL WORKS (ALL PROVISIONAL)', 47193.43552502702),
 ('DESIGN FEES', 194211.36000723258),
 ('PROJECT MANAGEMENT FEES', '=623262.5-175000'),
]
LET = list('ABCDEFGHJKLMNPQRS')
# helper columns (outside print area A:D): G = estimate before discount
gs['G5'] = 'ESTIMATE AMOUNT USD (BEFORE DISCOUNT) - REV 6'; gs['I5'] = 'CHECK'
for c in ('G5', 'I5'): gs[c].font = HELP_HDR; gs[c].alignment = Alignment(wrap_text=True)
rows = []
for i, (desc, amt) in enumerate(items):
    r = 7 + 2 * i; rows.append(r)
    gs[f'A{r}'] = LET[i]; gs[f'B{r}'] = desc
    gs[f'G{r}'] = amt; gs[f'G{r}'].fill = INFILL; gs[f'G{r}'].number_format = '#,##0.00'
    gs[f'D{r}'] = f'=G{r}*$G$52'
gs['D7'] = "='PRELIMS & GEN. ITEMS'!F41"
gs['I7'] = '=ROUND(D7-G7*$G$52,2)'; gs['I7'].number_format = '#,##0.00'
gs['F46'] = 'Estimate sub-total (before discount)'; gs['G46'] = '=SUM(G6:G43)'
gs['F48'] = 'Less: discount'; gs['G48'] = '=G46-G50'
gs['F50'] = 'Agreed discounted contract sum'; gs['G50'] = 6250000; gs['G50'].fill = INFILL
gs['F52'] = 'Discount factor (applied to every bill)'; gs['G52'] = '=G50/G46'
for r in (46, 48, 50):
    gs[f'G{r}'].number_format = '#,##0.00'
gs['G52'].number_format = '0.000000'
for r in (46, 48, 50, 52): gs[f'F{r}'].font = HELP_HDR
gs.column_dimensions['F'].width = 36; gs.column_dimensions['G'].width = 22; gs.column_dimensions['I'].width = 12
gs['B46'] = 'SUB - TOTAL (1)'; gs['C46'] = 'US $'; gs['D46'] = '=SUM(D6:D43)'
gs['B48'] = 'ADD: CONTINGENCY'
gs['B51'] = 'SUB - TOTAL (2)'; gs['C51'] = 'US $'; gs['D51'] = '=D48+D46'
gs['B53'] = 'NOTE: TAXES (NHIL, GETFUND LEVY, COVID LEVY AND VAT) ARE NOT INCLUDED'
gs['B61'] = 'TOTAL COST OF THE WORKS (EXCLUDING TAXES)'; gs['C61'] = 'US $'; gs['D61'] = '=D51'
gs['B64'] = 'SIGNED: ……………………………………………………………'
gs['B67'] = 'DATE: ………………………………………...........………….'
gs['F61'] = 'Check: total = agreed sum'; gs['F61'].font = HELP_HDR; gs['G61'] = '=ROUND(D61-G50,2)'

# ---------------- PRELIMS & GEN. ITEMS ----------------
pr = wb['PRELIMS & GEN. ITEMS']
pr['A2'] = 'PROJECT: PROPOSED HEAD OFFICE FOR PRIORITY INSURANCE'
pr['A3'] = 'LOCATION: RIDGE, ACCRA'
pr['A4'] = 'CLIENT: PRIORITY INSURANCE'
pr['B48'] = 'PRIORITY INSURANCE COMPANY LIMITED'
pr['B60'] = pr['B60'].value.replace('The Site is located at: Accra in the Greater Accra Region',
                                    'The Site is located at: Ridge, Accra in the Greater Accra Region')
pr['B102'] = 'KITCHEN CABINETRY/WALL CLADDING'
priced = [112, 115, 118, 124, 136, 139, 148, 151, 165, 189, 205, 215, 218, 221, 230, 233, 236, 239, 249, 252, 264, 267, 270, 302, 305, 337, 340]
BAL = 189   # Plant, Tools and Vehicles carries the rounding balance
# helper block (outside print area A:F)
pr['J8'] = 'THOROUGHBRED PLACE AMOUNT USD (REFERENCE)'; pr['K8'] = 'PRO-RATA AMOUNT USD (ROUNDED)'
for c in ('J8', 'K8'): pr[c].font = HELP_HDR; pr[c].alignment = Alignment(wrap_text=True, vertical='center')
pr['J10'] = 'Discounted Prelims amount (GEN SUMMARY Bill 1)'; pr['K10'] = "='GEN SUMMARY'!G7*'GEN SUMMARY'!G52"
pr['J11'] = 'Thoroughbred Place Prelims total'; pr['K11'] = '=SUM(J108:J370)'
pr['J12'] = 'Pro-rata factor'; pr['K12'] = '=K10/K11'
pr['J13'] = 'Rounding balance (added to Plant, Tools & Vehicles)'; pr['K13'] = '=K10-SUM(K108:K370)'
pr['J14'] = 'Check: Prelims total less discounted amount'; pr['K14'] = '=ROUND(F41-K10,2)'
for r in range(10, 15):
    pr[f'J{r}'].font = HELP_HDR; pr[f'J{r}'].alignment = Alignment(wrap_text=True)
    pr[f'K{r}'].number_format = '#,##0.00'
pr['K12'].number_format = '0.000000'
for r in priced:
    old = pr[f'F{r}'].value
    ref_amt = eval(old[1:]) if isinstance(old, str) and old.startswith('=') else old
    pr[f'J{r}'] = old if isinstance(old, str) else ref_amt
    pr[f'J{r}'].fill = INFILL; pr[f'J{r}'].number_format = '#,##0.00'
    pr[f'K{r}'] = f'=ROUND(J{r}*$K$12,2)'; pr[f'K{r}'].number_format = '#,##0.00'
    pr[f'F{r}'] = f'=K{r}+$K$13' if r == BAL else f'=K{r}'
pr.column_dimensions['J'].width = 30; pr.column_dimensions['K'].width = 22

wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(out)
print('saved', out)
