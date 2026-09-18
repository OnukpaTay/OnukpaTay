import openpyxl, json, datetime
from collections import OrderedDict

wb = openpyxl.load_workbook('source.xlsx', data_only=True)
ws = wb['18-09-2025']

def fmt(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%d %b %Y')
    return str(v).strip() if v else ''

rows=[]
proj=owner=None
for r in range(5,27):
    p = ws.cell(r,1).value or proj
    proj = p
    o = ws.cell(r,6).value or owner
    owner = o
    rows.append({
        'row': r,
        'project': str(p).strip(),
        'deliverable': str(ws.cell(r,2).value or '').strip(),
        'due_raw': ws.cell(r,3).value,
        'due': fmt(ws.cell(r,3).value),
        'planned': float(ws.cell(r,4).value or 0),
        'actual': float(ws.cell(r,5).value or 0),
        'owners': [x.strip() for x in str(o or '').split('\n') if x.strip()],
        'remark': str(ws.cell(r,7).value or '').strip(),
    })

n=len(rows)
done=[r for r in rows if r['actual']>=1]
prog=[r for r in rows if 0<r['actual']<1]
notst=[r for r in rows if r['actual']==0]
sp=sum(r['planned'] for r in rows); sa=sum(r['actual'] for r in rows)

print(f"TOTAL deliverables: {n}")
print(f"Projects: {len(set(r['project'] for r in rows))}")
print(f"Delivered: {len(done)}  In progress: {len(prog)}  Not started: {len(notst)}")
print(f"Planned units {sp} ({sp/n*100:.1f}%)  Actual units {sa} ({sa/n*100:.1f}%)  variance {sa-sp:+.1f} units ({(sa-sp)/n*100:+.1f} pts)")
print()
print("--- OUTSTANDING (actual < 1) ---")
for r in sorted(prog+notst, key=lambda x: x['actual']):
    print(f"  {r['project']:32s} | {r['deliverable'][:38]:38s} | due {r['due']:12s} | P {r['planned']*100:5.0f}% A {r['actual']*100:5.0f}% | gap {(r['planned']-r['actual'])*100:5.0f} | {'/'.join(r['owners'])} | {r['remark'][:50]}")
print()
print("--- BY PROJECT ---")
byp=OrderedDict()
for r in rows:
    d=byp.setdefault(r['project'], {'n':0,'p':0.0,'a':0.0,'owners':set()})
    d['n']+=1; d['p']+=r['planned']; d['a']+=r['actual']; d['owners'].update(r['owners'])
for k,v in byp.items():
    print(f"  {k:32s} n={v['n']}  planned {v['p']/v['n']*100:5.1f}%  actual {v['a']/v['n']*100:5.1f}%  owners={','.join(sorted(v['owners']))}")
print()
print("--- BY OWNER (shared items counted for each named person) ---")
byo={}
for r in rows:
    for o in r['owners']:
        d=byo.setdefault(o, {'n':0,'done':0})
        d['n']+=1
        if r['actual']>=1: d['done']+=1
for k,v in sorted(byo.items(), key=lambda x:-x[1]['n']):
    print(f"  {k:10s} {v['n']:2d} deliverables, {v['done']:2d} delivered")
print()
print("--- REMARKS ---")
for r in rows:
    if r['remark']: print(f"  {r['project']} / {r['deliverable']}: {r['remark']}")
json.dump({'rows':rows,'byproject':{k:{'n':v['n'],'p':v['p'],'a':v['a'],'owners':sorted(v['owners'])} for k,v in byp.items()},'byowner':byo}, open('data.json','w'), default=str, indent=1)
